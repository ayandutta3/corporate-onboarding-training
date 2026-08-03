import os
import aiofiles
from fastapi import UploadFile
import pymupdf
import docx
from pptx import Presentation
from langchain_openai import ChatOpenAI
from langchain_core.prompts import PromptTemplate
from app.configuration.settings import get_settings
from app.configuration.http_client import get_http_client, get_async_http_client

settings = get_settings()

async def save_upload_file(upload_file: UploadFile, destination_path: str) -> str:
    os.makedirs(os.path.dirname(destination_path), exist_ok=True)
    async with aiofiles.open(destination_path, 'wb') as out_file:
        while content := await upload_file.read(1024 * 1024):
            await out_file.write(content)
    return destination_path

async def extract_text_from_file(file_path: str, filename: str) -> str:
    ext = filename.split('.')[-1].lower()
    text = ""
    
    try:
        if ext == 'pdf':
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text()
        elif ext == 'docx':
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == 'pptx':
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext in ['png', 'jpg', 'jpeg']:
            try:
                from paddleocr import PaddleOCR
                ocr = PaddleOCR(use_angle_cls=True, lang='en')
                result = ocr.ocr(file_path, cls=True)
                if result and result[0]:
                    for line in result[0]:
                        text += line[1][0] + " "
            except Exception as e:
                print(f"OCR Extraction Warning: {e}")
                text = f"Image Document: {filename}"
        elif ext in ['mp4', 'avi', 'mov', 'm4v', 'mkv', 'mp3', 'wav', 'm4a', 'flac', 'ogg', 'aac']:
            import uuid
            temp_audio_file = None
            target_audio_path = file_path

            # For Video files: extract audio track to temporary .mp3 using MoviePy
            if ext in ['mp4', 'avi', 'mov', 'm4v', 'mkv']:
                temp_audio_file = f"uploads/temp_audio_{uuid.uuid4().hex}.mp3"
                try:
                    from moviepy.editor import VideoFileClip
                except ImportError:
                    try:
                        from moviepy import VideoFileClip
                    except ImportError:
                        VideoFileClip = None

                if VideoFileClip:
                    try:
                        clip = VideoFileClip(file_path)
                        if clip.audio:
                            clip.audio.write_audiofile(temp_audio_file, verbose=False, logger=None)
                            clip.close()
                            target_audio_path = temp_audio_file
                        else:
                            clip.close()
                    except Exception as clip_err:
                        print(f"MoviePy audio extraction warning for {filename}: {clip_err}")

            try:
                from openai import AsyncOpenAI
                client = AsyncOpenAI(
                    api_key=settings.openai_api_key,
                    base_url=settings.openai_api_base,
                    http_client=get_async_http_client()
                )
                
                with open(target_audio_path, "rb") as audio_file:
                    transcription = await client.audio.transcriptions.create(
                        model=settings.whisper_model,
                        file=audio_file,
                        response_format="verbose_json"
                    )
                
                segments = getattr(transcription, "segments", None)
                if not segments and isinstance(transcription, dict):
                    segments = transcription.get("segments")

                if segments:
                    text = group_segments_into_paragraphs(segments)
                else:
                    text = getattr(transcription, "text", str(transcription))
            except Exception as e:
                print(f"Whisper API transcription error for {filename}: {e}")
                # Fallback to local whisper model if available
                try:
                    import whisper
                    model = whisper.load_model("base")
                    result = model.transcribe(target_audio_path)
                    segments = result.get("segments", [])
                    if segments:
                        text = group_segments_into_paragraphs(segments)
                    else:
                        text = result.get("text", "")
                except Exception as local_e:
                    print(f"Local Whisper fallback error for {filename}: {local_e}")
                    text = f"Media File ({filename}): Speech audio content ingested."
            finally:
                if temp_audio_file and os.path.exists(temp_audio_file):
                    try:
                        os.remove(temp_audio_file)
                    except Exception:
                        pass
        elif ext == 'txt':
            async with aiofiles.open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = await f.read()
    except Exception as general_err:
        print(f"Text extraction error for {filename}: {general_err}")
        text = f"Document content for {filename}"
    
    return text.strip()

def format_seconds_to_timestamp(seconds: float) -> str:
    secs = int(max(0, seconds))
    hrs = secs // 3600
    mins = (secs % 3600) // 60
    rem_secs = secs % 60
    if hrs > 0:
        return f"{hrs:02d}:{mins:02d}:{rem_secs:02d}"
    return f"{mins:02d}:{rem_secs:02d}"

def group_segments_into_paragraphs(segments: list, max_block_duration_sec: float = 30.0, max_block_chars: int = 350) -> str:
    if not segments:
        return ""
    
    paragraph_blocks = []
    curr_start = None
    curr_end = None
    curr_texts = []
    curr_chars = 0

    for seg in segments:
        if isinstance(seg, dict):
            s_start = float(seg.get("start", 0))
            s_end = float(seg.get("end", 0))
            s_text = seg.get("text", "").strip()
        else:
            s_start = float(getattr(seg, "start", 0))
            s_end = float(getattr(seg, "end", 0))
            s_text = getattr(seg, "text", "").strip()

        if not s_text:
            continue

        if curr_start is None:
            curr_start = s_start
            curr_end = s_end
            curr_texts = [s_text]
            curr_chars = len(s_text)
        else:
            duration = s_end - curr_start
            if duration >= max_block_duration_sec or (curr_chars + len(s_text)) >= max_block_chars:
                start_ts = format_seconds_to_timestamp(curr_start)
                end_ts = format_seconds_to_timestamp(curr_end)
                combined_text = " ".join(curr_texts)
                paragraph_blocks.append(f"[{start_ts} - {end_ts}] {combined_text}")

                curr_start = s_start
                curr_end = s_end
                curr_texts = [s_text]
                curr_chars = len(s_text)
            else:
                curr_end = s_end
                curr_texts.append(s_text)
                curr_chars += len(s_text) + 1

    if curr_texts and curr_start is not None:
        start_ts = format_seconds_to_timestamp(curr_start)
        end_ts = format_seconds_to_timestamp(curr_end)
        combined_text = " ".join(curr_texts)
        paragraph_blocks.append(f"[{start_ts} - {end_ts}] {combined_text}")

    return "\n\n".join(paragraph_blocks)

async def generate_ai_summary_and_tags(text: str):
    if not text:
        return "No text content available for summary.", ["document", "ingested"]
    
    try:
        llm = ChatOpenAI(
            model=settings.llm_model,
            temperature=0,
            openai_api_key=settings.openai_api_key,
            openai_api_base=settings.openai_api_base,
            http_client=get_http_client(),
            http_async_client=get_async_http_client()
        )
        
        prompt = PromptTemplate(
            input_variables=["text"],
            template="Analyze the following text. Provide a brief summary and a list of 3-5 tags separated by commas.\n\nText: {text}\n\nFormat: Summary: [summary]\nTags: [tag1, tag2, tag3]"
        )
        
        chain = prompt | llm
        result = await chain.ainvoke({"text": text[:3000]}) # Limit text length for summary
        content = result.content
        
        content_lower = content.lower()
        summary = ""
        tags = []
        
        summary_idx = content_lower.find("summary:")
        tags_idx = content_lower.find("tags:")
        
        if summary_idx != -1 and tags_idx != -1:
            if summary_idx < tags_idx:
                summary_part = content[summary_idx + 8 : tags_idx].strip()
                tags_part = content[tags_idx + 5 :].strip()
            else:
                tags_part = content[tags_idx + 5 : summary_idx].strip()
                summary_part = content[summary_idx + 8 :].strip()
                
            summary = summary_part.strip("[]")
            tags = [t.strip("[]").strip() for t in tags_part.split(",") if t.strip()]
        else:
            lines = [line.strip() for line in content.split("\n") if line.strip()]
            for line in lines:
                if line.lower().startswith("summary:"):
                    summary = line[8:].strip("[]").strip()
                elif line.lower().startswith("tags:"):
                    tags = [t.strip("[]").strip() for t in line[5:].split(",") if t.strip()]
            
            if not summary and lines:
                summary = lines[0]
            if not tags and len(lines) > 1:
                tags = [t.strip("[]").strip() for t in lines[1].split(",") if t.strip()]
                
        if not summary:
            summary = text[:200]
        if not tags:
            tags = ["document", "onboarding"]
            
        return summary, tags
    except Exception as e:
        print(f"AI Summary generation fallback due to: {e}")
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        summary = lines[0][:200] if lines else "Document content ingested successfully."
        tags = ["document", "ingested"]
        return summary, tags
